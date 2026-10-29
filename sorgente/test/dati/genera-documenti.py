# Crea i documenti di prova Excel, Word e vCard. Uso: python3 test/dati/genera-documenti.py
import os, io
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImg
from openpyxl.styles import numbers
from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw
D=os.path.join(os.path.dirname(os.path.abspath(__file__)),'documenti')
def foto(col):
    im=Image.new('RGB',(160,160),col); d=ImageDraw.Draw(im); d.ellipse((42,42,118,118),fill='white'); b=io.BytesIO(); im.save(b,'PNG'); b.seek(0); return b
# --- Excel: materiale per impianto domestico, colonna Marca, sconto in percentuale, foto ancorate
wb=Workbook(); ws=wb.active; ws.title='Offerta'
ws['A1']='MATERIALE ELETTRICO ROSSI & FIGLI S.n.c.'; ws['A2']='Via Nazionale 101 - 75012 Bernalda (MT) - P.IVA 01234567897 - tel 0835 543210'
ws['A3']='Spett.le IECI S.r.l. - Via dei Pesci 82/D Matera'; ws['A4']='Preventivo n. 77 del 25/10/2026'; ws['A5']='Riferimento: Verdi Anna - Via Dante 10, 75100 Matera - tel 347 1234567'
H=['Codice','Descrizione','Marca','UM','Quantità','Prezzo listino','Sconto','Prezzo netto','Totale']; ws.append([]); ws.append(H)
R=[('14001','Interruttore 1P 16AX bianco Plana','Vimar','NR',10,7.46,0.35,4.85,48.50,'#c62828'),
   ('14201','Presa 2P+T 16A bipasso Plana bianco','Vimar','NR',12,9.23,0.35,6.00,72.00,'#2e7d32'),
   ('14653','Placca 3 moduli Plana bianco','Vimar','NR',15,3.85,0.35,2.50,37.50,'#1565c0'),
   ('V71303','Scatola da incasso 3 moduli','Vimar','NR',15,0.92,0.35,0.60,9.00,None),
   ('GW40004','Centralino da incasso 12 moduli IP40','Gewiss','NR',1,58.46,0.35,38.00,38.00,None),
   ('FS17-1.5','Cavo FS17 1x1,5 mmq nero (matassa 100 m)','','NR',3,30.77,0.35,20.00,60.00,None)]
for i,r in enumerate(R):
    ws.append(list(r[:9])); row=ws.max_row; ws.cell(row=row,column=7).number_format='0%'
    for cc in (6,8,9): ws.cell(row=row,column=cc).number_format='#,##0.00'
    if r[9]: img=XImg(foto(r[9])); img.width=40; img.height=40; ws.add_image(img,'J%d'%row)
ws.append([]);
for lab,v in (('Totale imponibile',265.00),('IVA 22%',58.30),('Totale',323.30)):
    ws.append(['',lab,'','','','','','',v]); ws.cell(row=ws.max_row,column=9).number_format='#,##0.00'
wb.save(os.path.join(D,'fornitore-domestico.xlsx'))
# --- Word: impianto di allarme, foto nelle celle della tabella
doc=Document(); doc.add_paragraph('SICUREZZA LUCANA S.r.l. - Via Appia 21 - 85100 Potenza - P.IVA 09876543217 - info@sicurezzalucana.it')
doc.add_paragraph('Alla cortese attenzione di IECI S.r.l.'); doc.add_paragraph('Offerta n. 2026-311 del 24/10/2026')
doc.add_paragraph('Cantiere: Condominio Le Rose, Via Gramsci 8, 75100 Matera')
A=[('38181','Ajax Hub 2 Plus centrale di allarme bianca','NR','1','289,00','289,00','#ef6c00'),('38184','Ajax MotionProtect sensore volumetrico PIR bianco','NR','6','52,00','312,00','#00838f'),('38188','Ajax DoorProtect contatto magnetico bianco','NR','8','29,00','232,00','#5d4037'),('38190','Ajax HomeSiren sirena da interno','NR','1','58,00','58,00',None),('38192','Ajax KeyPad tastiera wireless','NR','1','79,00','79,00',None)]
t=doc.add_table(rows=1,cols=7); hd=['Codice','Foto','Descrizione','UM','Q.tà','Prezzo','Importo']
for i,h in enumerate(hd): t.rows[0].cells[i].text=h
for a in A:
    c=t.add_row().cells; c[0].text=a[0]; c[2].text=a[1]; c[3].text=a[2]; c[4].text=a[3]; c[5].text=a[4]; c[6].text=a[5]
    if a[6]: c[1].paragraphs[0].add_run().add_picture(foto(a[6]),width=Cm(1.2))
doc.add_paragraph('Totale imponibile 970,00'); doc.add_paragraph('IVA 22% 213,40'); doc.add_paragraph('Totale offerta 1.183,40')
doc.save(os.path.join(D,'fornitore-allarme.docx'))
# --- vCard dal telefono (indirizzo e nota codificati)
open(os.path.join(D,'cliente.vcf'),'w').write('BEGIN:VCARD\r\nVERSION:2.1\r\nN:Verdi;Anna;;;\r\nFN:Anna Verdi\r\nTEL;CELL:+39 347 123 4567\r\nTEL;HOME:0835 111222\r\nEMAIL;HOME:anna.verdi@gmail.com\r\nADR;HOME;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:;;Via San Rocco 5;Matera;MT;75100;Italia\r\nNOTE;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:C.F. VRDNNA80A41F052D =E2=80=93 cliente dal 2024\r\nEND:VCARD\r\n')
print('ok')
